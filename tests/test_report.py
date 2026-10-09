from datetime import date

from openpyxl import load_workbook

from sap_report.main import genera_report
from sap_report.sorgenti.mock import SorgenteMock

GIORNO = date(2026, 10, 8)


def test_report_excel_ha_i_fogli_attesi(tmp_path):
    destinazione = genera_report(SorgenteMock(), GIORNO, tmp_path)
    assert destinazione.name == "report_ordini_2026-10-08.xlsx"

    wb = load_workbook(destinazione)
    assert wb.sheetnames == [
        "Riepilogo",
        "Per cliente",
        "Per stato",
        "Ordini in ritardo",
        "Dettaglio ordini",
    ]
